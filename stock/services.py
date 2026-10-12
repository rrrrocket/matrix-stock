from collections import defaultdict
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from datetime import datetime
from io import BytesIO, StringIO
import csv
from uuid import UUID, uuid4
from zipfile import BadZipFile
from urllib.parse import urlsplit

from openpyxl import load_workbook
from django.db import transaction
from django.db.models import F, Sum
from django.utils.dateparse import parse_datetime

from .models import ExternalOrder, OrderAlert, OrderLine, OutboundDocument, OutboundDocumentLine, StockBalance, StockItem, StockMovement, Warehouse


class StockError(ValueError):
    pass


def _positive_int(value, label):
    if isinstance(value, bool):
        raise StockError(f"{label}必须是正整数")
    try:
        number = int(value)
    except (TypeError, ValueError):
        raise StockError(f"{label}必须是正整数") from None
    if number <= 0 or number > 2_147_483_647 or str(value).strip() != str(number):
        raise StockError(f"{label}必须是正整数")
    return number


def normalize_sku(value):
    sku = str(value or "").strip()
    if not sku or len(sku) > 150:
        raise StockError("货号不能为空且不能超过150字符")
    return sku


def parse_cost(value, *, required=False):
    if value is None or str(value).strip() == "":
        if required:
            raise StockError("请先填写库存成本单价")
        return None
    try:
        cost = Decimal(str(value).strip())
    except (InvalidOperation, ValueError):
        raise StockError("成本单价格式无效") from None
    if not cost.is_finite() or cost < 0 or cost >= Decimal("10000000000") or cost.as_tuple().exponent < -4:
        raise StockError("成本单价须为非负数，最多四位小数")
    return cost.quantize(Decimal("0.0001"))


def _money(value):
    return value.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


@transaction.atomic
def record_inbound(*, warehouse_id, sku, quantity, product_name="", note="", operator=None, movement_type=StockMovement.Type.INBOUND, unit_cost=None):
    warehouse_id = _positive_int(warehouse_id, "仓库ID")
    quantity = _positive_int(quantity, "入库数量")
    sku = normalize_sku(sku)
    warehouse = Warehouse.objects.get(pk=warehouse_id, owner=operator, is_active=True)
    item, _ = StockItem.objects.get_or_create(owner=operator, sku=sku, defaults={"name": str(product_name or "").strip()[:255]})
    if product_name and not item.name:
        item.name = str(product_name).strip()[:255]
        item.save(update_fields=["name"])
    balance, _ = StockBalance.objects.select_for_update().get_or_create(warehouse=warehouse, item=item)
    incoming_cost = parse_cost(unit_cost) if unit_cost is not None else item.reference_cost
    if incoming_cost is not None:
        if unit_cost is not None and str(unit_cost).strip() and item.reference_cost != incoming_cost:
            item.reference_cost = incoming_cost
            item.save(update_fields=["reference_cost"])
        if balance.on_hand and balance.avg_unit_cost is None:
            raise StockError("原库存尚未定价，请先设置库存成本单价")
        old_value = (balance.avg_unit_cost or Decimal(0)) * balance.on_hand
        balance.avg_unit_cost = (old_value + incoming_cost * quantity) / (balance.on_hand + quantity)
        balance.avg_unit_cost = balance.avg_unit_cost.quantize(Decimal("0.0001"), rounding=ROUND_HALF_UP)
        balance.save(update_fields=["avg_unit_cost"])
    StockBalance.objects.filter(pk=balance.pk).update(on_hand=F("on_hand") + quantity)
    balance.refresh_from_db()
    movement = StockMovement.objects.create(
        warehouse=warehouse, item=item, movement_type=movement_type,
        quantity_change=quantity, before_quantity=balance.on_hand - quantity,
        after_quantity=balance.on_hand, note=str(note or "")[:500], operator=operator,
    )
    for order_id in OrderLine.objects.filter(
        sku=sku, order__owner=operator, is_present=True, is_skipped=False, order__status=ExternalOrder.Status.OPEN,
    ).values_list("order_id", flat=True).distinct():
        sync_alert_for_order(ExternalOrder.objects.get(pk=order_id))
    return movement


@transaction.atomic
def record_batch_inbound(*, warehouse_id, rows, default_quantity, operator):
    warehouse_id = _positive_int(warehouse_id, "仓库ID")
    if not Warehouse.objects.filter(pk=warehouse_id, owner=operator, is_active=True).exists():
        raise StockError("仓库不存在")
    default_quantity = _positive_int(default_quantity, "默认入库数量")
    if not isinstance(rows, list) or not rows or len(rows) > 100:
        raise StockError("每批请输入 1 至 100 条货号")
    parsed = []
    for row in rows:
        sku = normalize_sku(row.get("sku"))
        quantity = _positive_int(row.get("quantity") or default_quantity, "入库数量")
        note = str(row.get("note") or "").strip()
        if len(note) > 500:
            raise StockError("备注不能超过 500 字符")
        parsed.append((sku, quantity, note, row.get("unit_cost")))
    return [
        record_inbound(warehouse_id=warehouse_id, sku=sku, quantity=quantity, note=note, operator=operator, unit_cost=unit_cost)
        for sku, quantity, note, unit_cost in parsed
    ]


@transaction.atomic
def adjust_stock(*, warehouse_id, sku, new_quantity, note, operator=None, expected_quantity=None,
                 movement_type=StockMovement.Type.ADJUST):
    warehouse_id = _positive_int(warehouse_id, "仓库ID")
    if isinstance(new_quantity, bool):
        raise StockError("盘点数量必须是非负整数")
    raw_quantity = new_quantity
    try:
        new_quantity = int(raw_quantity)
    except (TypeError, ValueError):
        raise StockError("盘点数量必须是非负整数") from None
    if new_quantity < 0 or new_quantity > 2_147_483_647 or str(raw_quantity).strip() != str(new_quantity):
        raise StockError("盘点数量必须是非负整数")
    if not str(note or "").strip():
        raise StockError("盘点调整必须填写原因")
    warehouse = Warehouse.objects.get(pk=warehouse_id, owner=operator, is_active=True)
    item = StockItem.objects.get(owner=operator, sku=normalize_sku(sku))
    balance, _ = StockBalance.objects.select_for_update().get_or_create(warehouse=warehouse, item=item)
    before = balance.on_hand
    if expected_quantity is not None:
        try:
            expected = int(expected_quantity)
        except (TypeError, ValueError):
            raise StockError("无效的原库存数量") from None
        if expected < 0 or str(expected_quantity).strip() != str(expected):
            raise StockError("无效的原库存数量")
        if before != expected:
            raise StockError("库存已变化，请刷新后重试")
    if before == new_quantity:
        raise StockError("新库存与当前库存相同")
    balance.on_hand = new_quantity
    balance.save(update_fields=["on_hand", "updated_at"])
    return StockMovement.objects.create(
        warehouse=warehouse, item=item, movement_type=movement_type,
        quantity_change=new_quantity - before, before_quantity=before,
        after_quantity=new_quantity, note=str(note).strip()[:500], operator=operator,
    )


@transaction.atomic
def record_outbound(*, line_id, warehouse_id, quantity, tracking_number, operator, operation_id):
    warehouse_id = _positive_int(warehouse_id, "仓库ID")
    quantity = _positive_int(quantity, "出库数量")
    tracking_number = str(tracking_number or "").strip()
    if not tracking_number or len(tracking_number) > 200 or any(ord(char) < 32 for char in tracking_number):
        raise StockError("请填写有效的国内快递单号（最多200字符）")
    try:
        operation_id = UUID(str(operation_id))
    except (TypeError, ValueError, AttributeError):
        raise StockError("无效的操作编号") from None
    previous = StockMovement.objects.filter(operation_id=operation_id).first()
    if previous:
        if (
            previous.movement_type != StockMovement.Type.OUTBOUND
            or previous.order_line.order.owner_id != operator.pk
            or previous.order_line_id != line_id
            or str(previous.warehouse_id) != str(warehouse_id)
            or previous.quantity_change != -quantity
            or previous.tracking_number != tracking_number
        ):
            raise StockError("操作编号已用于其他业务")
        return previous

    line = OrderLine.objects.select_for_update().select_related("order").get(pk=line_id, order__owner=operator)
    if line.order.status != ExternalOrder.Status.OPEN or not line.is_present:
        raise StockError("订单已取消、结束或商品已不在订单中")
    if line.is_skipped:
        raise StockError("该订单商品已设为暂不处理")
    if quantity > line.remaining_quantity:
        raise StockError(f"本单最多还能出库 {line.remaining_quantity} 件")
    warehouse = Warehouse.objects.get(pk=warehouse_id, owner=operator, is_active=True)
    item = StockItem.objects.filter(owner=operator, sku=line.sku).first()
    if not item:
        raise StockError("库存系统中找不到该货号")
    balance = StockBalance.objects.filter(warehouse=warehouse, item=item).first()
    if not balance:
        raise StockError("所选仓库没有该货号库存")
    updated = StockBalance.objects.filter(pk=balance.pk, on_hand__gte=quantity).update(on_hand=F("on_hand") - quantity)
    if not updated:
        raise StockError("仓库库存不足，请刷新页面")
    updated_line = OrderLine.objects.filter(pk=line.pk, quantity__gte=F("outbound_quantity") + quantity).update(
        outbound_quantity=F("outbound_quantity") + quantity
    )
    if not updated_line:
        raise StockError("订单剩余可出库数量已变化，请刷新页面")
    balance.refresh_from_db()
    return StockMovement.objects.create(
        operation_id=operation_id, warehouse=warehouse, item=item, order_line=line,
        tracking_number=tracking_number,
        movement_type=StockMovement.Type.OUTBOUND, quantity_change=-quantity,
        before_quantity=balance.on_hand + quantity, after_quantity=balance.on_hand,
        note=f"{line.order.platform} / {line.order.store_name} / {line.order.external_order_id}",
        operator=operator,
    )


def _new_document(*, owner, warehouse, source, operator, purpose="", recipient="", note="", request_id=None):
    document = OutboundDocument.objects.create(
        owner=owner, warehouse=warehouse, source=source, operator=operator,
        purpose=purpose, recipient=recipient, note=note, document_no=uuid4().hex,
        **({"request_id": request_id} if request_id else {}),
    )
    document.document_no = f"CK{document.created_at:%Y%m%d}-{document.pk:06d}"
    document.save(update_fields=["document_no"])
    return document


def _attach_document_line(document, movement, unit_cost):
    quantity = -movement.quantity_change
    amount = _money(unit_cost * quantity) if unit_cost is not None else None
    if amount is not None and (amount >= Decimal("10000000000000000") or
                               (document.total_amount is not None and document.total_amount + amount >= Decimal("10000000000000000"))):
        raise StockError("出库金额超出允许范围")
    OutboundDocumentLine.objects.create(
        document=document, movement=movement, order_line=movement.order_line,
        sku=movement.item.sku, quantity=quantity, unit_cost=unit_cost,
        amount=amount, tracking_number=movement.tracking_number,
    )
    document.total_quantity += quantity
    if amount is None:
        document.total_amount = None
    elif document.total_amount is not None:
        document.total_amount += amount


@transaction.atomic
def record_order_outbound_document(*, line_id, warehouse_id, quantity, tracking_number, operator, operation_id):
    line_id = _positive_int(line_id, "订单商品ID")
    try:
        operation_id = UUID(str(operation_id))
    except (ValueError, TypeError, AttributeError):
        raise StockError("无效的操作编号") from None
    previous = StockMovement.objects.filter(operation_id=operation_id).first()
    if previous:
        movement = record_outbound(line_id=line_id, warehouse_id=warehouse_id, quantity=quantity,
                                   tracking_number=tracking_number, operator=operator, operation_id=operation_id)
        return movement.document_line.document, movement
    line = OrderLine.objects.select_related("order").get(pk=line_id, order__owner=operator)
    balance = StockBalance.objects.select_for_update().filter(
        warehouse_id=warehouse_id, warehouse__owner=operator, item__owner=operator, item__sku=line.sku,
    ).first()
    if not balance:
        raise StockError("所选仓库没有该货号库存")
    cost = parse_cost(balance.avg_unit_cost)
    movement = record_outbound(line_id=line_id, warehouse_id=warehouse_id, quantity=quantity,
                               tracking_number=tracking_number, operator=operator, operation_id=operation_id)
    document = _new_document(owner=operator, warehouse=movement.warehouse,
                             source=OutboundDocument.Source.ORDER, operator=operator)
    _attach_document_line(document, movement, cost)
    document.save(update_fields=["total_quantity", "total_amount"])
    return document, movement


@transaction.atomic
def record_batch_order_outbound(*, rows, operator):
    if not isinstance(rows, list) or not 1 <= len(rows) <= 100:
        raise StockError("每次请选择 1 至 100 条订单商品")
    if any(not isinstance(row, dict) for row in rows):
        raise StockError("订单商品格式无效")
    if len({str(row.get("line_id")) for row in rows}) != len(rows):
        raise StockError("同一订单商品不能重复选择")
    try:
        operation_ids = [UUID(str(row.get("operation_id"))) for row in rows]
    except (ValueError, TypeError, AttributeError):
        raise StockError("无效的操作编号") from None
    if len(set(operation_ids)) != len(operation_ids):
        raise StockError("操作编号不能重复")
    previous = [StockMovement.objects.filter(operation_id=operation_id).first() for operation_id in operation_ids]
    if any(previous) and not all(previous):
        raise StockError("部分订单已出库，请刷新后重新选择")
    if all(previous):
        documents = {}
        for row, movement in zip(rows, previous):
            document, _ = record_order_outbound_document(
                line_id=row.get("line_id"), warehouse_id=row.get("warehouse_id"),
                quantity=row.get("quantity"), tracking_number=row.get("tracking_number"),
                operator=operator, operation_id=row.get("operation_id"))
            documents[document.pk] = document
        return list(documents.values()), previous
    movements = []
    documents = {}
    for row in rows:
        document, movement = record_order_outbound_document(
            line_id=row.get("line_id"), warehouse_id=row.get("warehouse_id"),
            quantity=row.get("quantity"), tracking_number=row.get("tracking_number"),
            operator=operator, operation_id=row.get("operation_id"),
        )
        # Merge this transaction's individual documents into one document per warehouse.
        if movement.warehouse_id in documents and documents[movement.warehouse_id].pk != document.pk:
            target = documents[movement.warehouse_id]
            line = movement.document_line
            line.document = target
            line.save(update_fields=["document"])
            target.total_quantity += document.total_quantity
            if target.total_amount is None or document.total_amount is None:
                target.total_amount = None
            else:
                target.total_amount += document.total_amount
            target.save(update_fields=["total_quantity", "total_amount"])
            document.delete()
        else:
            documents[movement.warehouse_id] = document
        movements.append(movement)
    return list(documents.values()), movements


@transaction.atomic
def record_manual_outbound(*, warehouse_id, rows, purpose, recipient, note, operator, request_id):
    try:
        request_id = UUID(str(request_id))
    except (ValueError, TypeError, AttributeError):
        raise StockError("无效的操作编号") from None
    warehouse = Warehouse.objects.select_for_update().get(
        pk=_positive_int(warehouse_id, "仓库ID"), owner=operator, is_active=True,
    )
    existing = OutboundDocument.objects.filter(request_id=request_id).first()
    if existing:
        if existing.owner_id != operator.pk or existing.source != OutboundDocument.Source.MANUAL:
            raise StockError("操作编号已用于其他业务")
        submitted = [(normalize_sku(row.get("sku")), _positive_int(row.get("quantity"), "出库数量")) for row in rows]
        stored = list(existing.lines.order_by("id").values_list("sku", "quantity"))
        if (str(existing.warehouse_id) != str(warehouse_id) or existing.purpose != str(purpose or "").strip()
                or existing.recipient != str(recipient or "").strip() or existing.note != str(note or "").strip()
                or submitted != stored):
            raise StockError("操作编号已用于其他业务")
        return existing
    purpose = str(purpose or "").strip()
    recipient = str(recipient or "").strip()
    note = str(note or "").strip()
    if not purpose or len(purpose) > 200 or len(recipient) > 200 or len(note) > 500:
        raise StockError("请填写有效的出库用途、收货方和备注")
    if not isinstance(rows, list) or not 1 <= len(rows) <= 100:
        raise StockError("每批请输入 1 至 100 条货号")
    parsed = [(normalize_sku(row.get("sku")), _positive_int(row.get("quantity"), "出库数量")) for row in rows]
    if len({sku for sku, _ in parsed}) != len(parsed):
        raise StockError("同一货号请合并为一行")
    document = _new_document(owner=operator, warehouse=warehouse, source=OutboundDocument.Source.MANUAL,
                             operator=operator, purpose=purpose, recipient=recipient, note=note, request_id=request_id)
    for sku, quantity in parsed:
        balance = StockBalance.objects.select_for_update().select_related("item").filter(
            warehouse=warehouse, item__owner=operator, item__sku=sku,
        ).first()
        if not balance or balance.on_hand < quantity:
            raise StockError(f"{sku} 库存不足，请刷新后重试")
        cost = parse_cost(balance.avg_unit_cost, required=True)
        before = balance.on_hand
        balance.on_hand -= quantity
        balance.save(update_fields=["on_hand", "updated_at"])
        movement = StockMovement.objects.create(
            warehouse=warehouse, item=balance.item, movement_type=StockMovement.Type.MANUAL_OUTBOUND,
            quantity_change=-quantity, before_quantity=before, after_quantity=balance.on_hand,
            note=note or purpose, operator=operator,
        )
        _attach_document_line(document, movement, cost)
    document.save(update_fields=["total_quantity", "total_amount"])
    return document


def _required_text(payload, key, max_length):
    value = str(payload.get(key) or "").strip()
    if not value or len(value) > max_length:
        raise StockError(f"{key}不能为空且不能超过{max_length}字符")
    return value


@transaction.atomic
def upsert_external_order(payload, *, owner_id):
    if not isinstance(payload, dict):
        raise StockError("订单数据必须是对象")
    version = payload.get("version")
    if version is not None and (type(version) is not int or version <= 0):
        raise StockError("version必须是正整数")
    platform = _required_text(payload, "platform", 40)
    store_external_id = _required_text(payload, "store_id", 80)
    store_name = _required_text(payload, "store_name", 150)
    external_order_id = _required_text(payload, "order_id", 150)
    status = payload.get("status")
    if status not in ExternalOrder.Status.values:
        raise StockError("status只允许open、cancelled或closed")
    rows = payload.get("items")
    if not isinstance(rows, list) or not rows:
        raise StockError("items必须是非空列表")
    quantities = defaultdict(int)
    images = {}
    for row in rows:
        if not isinstance(row, dict):
            raise StockError("订单商品格式无效")
        sku = normalize_sku(row.get("sku"))
        quantities[sku] += _positive_int(row.get("quantity"), "订单数量")
        image_url = str(row.get("image_url") or "").strip()
        if not image_url or len(image_url) > 1000 or any(ord(char) < 32 for char in image_url):
            continue
        try:
            parsed_image = urlsplit(image_url)
        except ValueError:
            continue
        if (parsed_image.scheme in {"http", "https"} and parsed_image.netloc
                and not parsed_image.username and not parsed_image.password):
            images.setdefault(sku, image_url)
    ordered_at = payload.get("ordered_at")
    if ordered_at:
        if not isinstance(ordered_at, str):
            raise StockError("ordered_at必须是ISO时间字符串")
        parsed = parse_datetime(ordered_at)
        if not isinstance(parsed, datetime) or parsed.tzinfo is None:
            raise StockError("ordered_at必须包含时区")
        ordered_at = parsed
    else:
        ordered_at = None

    order, created = ExternalOrder.objects.select_for_update().get_or_create(
        owner_id=owner_id, platform=platform, store_external_id=store_external_id, external_order_id=external_order_id,
        defaults={"store_name": store_name, "status": status, "ordered_at": ordered_at, "snapshot_version": version or 0},
    )
    if not created:
        if order.snapshot_version and version is None:
            raise StockError("该订单需要带version的快照")
        if version is not None and version < order.snapshot_version:
            raise StockError("过期的订单快照")
        if version is not None and version == order.snapshot_version:
            return order, False
        order.store_name = store_name
        order.status = status
        order.ordered_at = ordered_at or order.ordered_at
        order.snapshot_version = version or 0
        order.save(update_fields=["store_name", "status", "ordered_at", "snapshot_version", "updated_at"])
    OrderLine.objects.filter(order=order).exclude(sku__in=quantities).update(is_present=False)
    for sku, quantity in quantities.items():
        line, line_created = OrderLine.objects.get_or_create(
            order=order, sku=sku,
            defaults={"image_url": images.get(sku, ""), "quantity": quantity},
        )
        if not line_created:
            line.quantity = quantity
            line.image_url = images.get(sku, "")
            line.is_present = True
            line.save(update_fields=["quantity", "image_url", "is_present"])
    sync_alert_for_order(order)
    return order, created


def matching_order_lines(order):
    if order.status != ExternalOrder.Status.OPEN:
        return []
    lines = list(order.lines.filter(is_present=True, is_skipped=False))
    available = dict(
        StockBalance.objects.filter(warehouse__owner=order.owner, item__sku__in=[line.sku for line in lines], on_hand__gt=0, warehouse__is_active=True)
        .values("item__sku").annotate(total=Sum("on_hand"))
        .values_list("item__sku", "total")
    )
    return [(line, available.get(line.sku, 0)) for line in lines if line.remaining_quantity > 0 and available.get(line.sku, 0) > 0]


def sync_alert_for_order(order):
    matched = matching_order_lines(order)
    if not matched:
        return None
    fingerprint = ",".join(sorted(line.sku for line, _ in matched))[:500]
    alert, created = OrderAlert.objects.get_or_create(order=order, defaults={"fingerprint": fingerprint})
    if not created and alert.fingerprint != fingerprint:
        alert.fingerprint = fingerprint
        alert.status = OrderAlert.Status.PENDING
        alert.next_attempt_at = None
        alert.last_error = ""
        alert.save(update_fields=["fingerprint", "status", "next_attempt_at", "last_error"])
    return alert


@transaction.atomic
def import_opening_stock(uploaded_file, operator):
    """Import initial counts once per warehouse/SKU; never overwrite an existing ledger."""
    if uploaded_file.size > 2 * 1024 * 1024:
        raise StockError("文件不能超过2MB")
    filename = uploaded_file.name.lower()
    data = uploaded_file.read()
    try:
        if filename.endswith(".csv"):
            reader = csv.DictReader(StringIO(data.decode("utf-8-sig")))
            rows = list(reader)
        elif filename.endswith(".xlsx"):
            workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
            sheet = workbook.active
            values = sheet.iter_rows(values_only=True)
            headers = [str(value or "").strip() for value in next(values)]
            rows = [dict(zip(headers, row)) for row in values]
            workbook.close()
        else:
            raise StockError("只支持 .xlsx 或 .csv 文件")
    except (UnicodeDecodeError, StopIteration, ValueError, OSError, BadZipFile) as exc:
        raise StockError(f"文件无法读取：{exc}") from exc
    if not rows or len(rows) > 1000:
        raise StockError("文件应包含1至1000条数据")
    required = {"仓库编码", "货号", "数量"}
    if not required.issubset(rows[0]):
        raise StockError("表头必须包含：仓库编码、货号、数量；可选商品名称")
    parsed = []
    seen = set()
    warehouses = {warehouse.code: warehouse for warehouse in Warehouse.objects.filter(owner=operator, is_active=True)}
    for index, row in enumerate(rows, start=2):
        code = str(row.get("仓库编码") or "").strip()
        warehouse = warehouses.get(code)
        if not warehouse:
            raise StockError(f"第{index}行：找不到启用的仓库编码 {code}")
        sku = normalize_sku(row.get("货号"))
        raw_quantity = row.get("数量")
        if isinstance(raw_quantity, float) and raw_quantity.is_integer():
            raw_quantity = int(raw_quantity)
        if str(raw_quantity or "").strip() == "0":
            continue
        quantity = _positive_int(raw_quantity, f"第{index}行数量")
        key = (warehouse.id, sku)
        if key in seen:
            raise StockError(f"第{index}行：同一仓库货号在文件中重复")
        seen.add(key)
        parsed.append((warehouse, sku, quantity, str(row.get("商品名称") or "").strip()[:255]))
    if not parsed:
        raise StockError("文件中没有大于0的期初库存")
    for warehouse, sku, _, _ in parsed:
        if StockMovement.objects.filter(warehouse=warehouse, item__owner=operator, item__sku=sku).exists():
            raise StockError(f"{warehouse.name} / {sku} 已有库存流水，不能重复导入期初库存")
        if StockBalance.objects.filter(warehouse=warehouse, item__owner=operator, item__sku=sku, on_hand__gt=0).exists():
            raise StockError(f"{warehouse.name} / {sku} 已有库存，不能重复导入期初库存")
    for warehouse, sku, quantity, name in parsed:
        record_inbound(
            warehouse_id=warehouse.id, sku=sku, quantity=quantity,
            product_name=name, operator=operator,
            movement_type=StockMovement.Type.OPENING,
            note="Excel期初库存导入",
        )
    return len(parsed)
